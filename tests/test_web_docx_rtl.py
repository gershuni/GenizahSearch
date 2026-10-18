# -*- coding: utf-8 -*-
"""Website Word exports: Hebrew paragraphs read right to left from the RIGHT
margin (2026-10-17).

`web.export_service.set_paragraph_rtl` aligned "right" and appended `w:bidi`;
Word reads a right-to-left paragraph aligned "right" as flush LEFT, so every
Hebrew paragraph of every website Word export sat on the left margin
(rendered through Word). It now uses the shared helper: Hebrew -> `w:bidi`
and no alignment of its own; Latin-only -> left to right, unaligned, as in
the rest of these English documents. `set_run_rtl_font` marks a run
right-to-left only when it has no Latin letters: a Hebrew library name and a
Latin shelfmark in one right-to-left run scrambled the result header.
"""
from __future__ import annotations

from unittest.mock import MagicMock

import pytest

docx = pytest.importorskip("docx")

from docx.oxml.ns import qn  # noqa: E402

from web.export_service import (ExportService, add_hebrew_paragraph,  # noqa: E402
                                set_paragraph_rtl, set_run_rtl_font)

HEB = 'ברוך אתה יי אלהינו מלך העולם (הוא נותן לחם לכל בשר).'


def _is_hebrew(text):
    return any("א" <= ch <= "ת" for ch in text)


def _bidi(p):
    return p._p.pPr is not None and p._p.pPr.find(qn("w:bidi")) is not None


def _jc(p):
    return None if p._p.pPr is None else p._p.pPr.find(qn("w:jc"))


def _run_rtl(run):
    rpr = run._r.rPr
    return rpr is not None and rpr.find(qn("w:rtl")) is not None


def test_a_hebrew_paragraph_starts_at_the_right_margin():
    doc = docx.Document()
    add_hebrew_paragraph(doc, HEB)
    p = doc.paragraphs[0]
    assert _bidi(p)
    assert _jc(p) is None, 'right-to-left and aligned "right" is flush LEFT in Word'


def test_a_latin_only_paragraph_stays_where_an_english_document_puts_it():
    doc = docx.Document()
    p = doc.add_paragraph("1. Cambridge University Library, T-S 12.345 - Score: 9")
    set_paragraph_rtl(p)
    assert not _bidi(p)
    assert _jc(p) is None


def test_only_a_run_without_latin_letters_is_marked_right_to_left():
    doc = docx.Document()
    p = doc.add_paragraph()
    hebrew = p.add_run("סדור (קטעים).")
    mixed = p.add_run("1. ספריית אוניברסיטת קיימברידג׳, T-S 12.345")
    set_run_rtl_font(hebrew)
    set_run_rtl_font(mixed)
    assert _run_rtl(hebrew)
    assert not _run_rtl(mixed), (
        "a right-to-left run holding a Latin shelfmark scrambles the header")


@pytest.fixture
def service():
    mgr = MagicMock()
    mgr.get_meta_for_id.return_value = ("T-S 12.345", "סדור (קטעים).")
    mgr.get_library_for_id.return_value = "CUL"
    mgr.parse_full_id_components.return_value = {}
    return ExportService(meta_mgr=mgr)


def _hebrew_paragraphs(blob):
    import io
    doc = docx.Document(io.BytesIO(blob))
    return doc, [p for p in doc.paragraphs if _is_hebrew(p.text)]


@pytest.mark.parametrize("kind", ["search", "parallels", "browse"])
def test_every_hebrew_paragraph_of_each_word_export_starts_at_the_right(service, kind):
    if kind == "search":
        blob, _ = service.export_search_results_word(
            [{'display': {'shelfmark': 'T-S 12.345', 'title': 'סדור (קטעים).',
                          'id': '990001234560205171'},
              'snippet': 'ברוך *אתה* יי', 'full_text': HEB}], search_query='ברוך')
    elif kind == "parallels":
        blob, _ = service.export_parallels_word(
            [{'raw_header': '990001234560205171_IE1_P000001_FL1', 'score': 120,
              'source_ctx': 'ברוך *אתה* יי', 'text': HEB}], [], source_text=HEB)
    else:
        blob, _ = service.export_browse_word(
            {'shelfmark': 'T-S 12.345', 'library_name': 'Cambridge University Library',
             'title': 'סדור (קטעים).', 'sys_id': '990001234560205171',
             'text': HEB, 'p_num': 1, 'lang': 'en'})
    doc, hebrew = _hebrew_paragraphs(blob)
    # The heading carries the query text in Hebrew too; headings are the
    # document's own (English) chrome and are not laid out here.
    body = [p for p in hebrew if not p.style.name.startswith(("Heading", "Title"))]
    assert body, [p.text for p in doc.paragraphs]
    for p in body:
        assert _bidi(p), p.text
        assert _jc(p) is None, f'flush left in Word: {p.text!r}'
        for run in p.runs:
            if any("A" <= ch.upper() <= "Z" for ch in run.text):
                assert not _run_rtl(run), run.text
