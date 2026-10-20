# -*- coding: utf-8 -*-
"""Each visitor's text comes out in their own language, also after an await.

``web/translations.py`` used to keep the language in one module-level value
that every page render overwrote, so a page that drew after an await (a
deferred fill, a search's results) or a download built by an API route could
come out in whichever language the LAST visitor's page was rendered in
(tracker: "Web text drawn after an `await` can come out in another visitor's
language"). The language now belongs to the page, the task the page started,
or the request; see the precedence comment in ``web/translations.py``.
"""

from __future__ import annotations

import asyncio
import threading
from types import SimpleNamespace

import pytest
from nicegui.slot import Slot
from nicegui.storage import request_contextvar

import web.translations as translations
from web.translations import get_language, set_language, tr, using_language

# A string with a Hebrew translation, so tr() shows which language it used.
_KEY = 'Search'
assert translations.TRANSLATIONS.get(_KEY) not in (None, _KEY)
_HE = translations.TRANSLATIONS[_KEY]


class _Page:
    """Stands in for a NiceGUI client: what ``Slot.parent.client`` returns."""


def _enter_page(page: _Page):
    """Put this asyncio task inside ``page``, as ``with page_client:`` does."""
    slot = SimpleNamespace(parent=SimpleNamespace(client=page))
    Slot.get_stack().append(slot)


def _leave_page():
    Slot.get_stack().pop()
    Slot.prune_stack()


@pytest.fixture(autouse=True)
def _process_default():
    saved = translations._current_lang
    token = request_contextvar.set(None)
    yield
    request_contextvar.reset(token)
    translations._current_lang = saved


def test_two_pages_rendering_concurrently_keep_their_own_language():
    """The reported bug: visitor A's page awaits, visitor B's page renders,
    and A's page then draws -- in A's language, not B's."""
    seen: dict = {}

    async def render(name, lang, other_rendered, mine_rendered):
        page = _Page()
        _enter_page(page)
        try:
            set_language(lang)  # what create_layout() does
            mine_rendered.set()
            await other_rendered.wait()  # the other visitor's render runs here
            await asyncio.sleep(0)
            seen[name] = (get_language(), tr(_KEY))
        finally:
            _leave_page()

    async def main():
        a_done, b_done = asyncio.Event(), asyncio.Event()
        await asyncio.gather(
            render('a', 'en', b_done, a_done),
            render('b', 'he', a_done, b_done),
        )

    asyncio.run(main())
    assert seen == {'a': ('en', _KEY), 'b': ('he', _HE)}


def test_an_event_handler_uses_its_page_language_not_the_last_render():
    """A click handler runs in a fresh task inside its page's slot."""
    seen = []

    async def main():
        english_page, hebrew_page = _Page(), _Page()

        async def build(page, lang):
            _enter_page(page)
            try:
                set_language(lang)
            finally:
                _leave_page()

        # Each page is built in its own request task; an event arrives in a
        # task of its own, from the socket's context, not from either build.
        await asyncio.create_task(build(english_page, 'en'))
        await asyncio.create_task(build(hebrew_page, 'he'))  # the last render was Hebrew

        async def click():  # a new task, as NiceGUI runs each event
            _enter_page(english_page)
            try:
                seen.append(tr(_KEY))
            finally:
                _leave_page()

        await asyncio.create_task(click())

    asyncio.run(main())
    assert seen == [_KEY]


def test_a_task_started_by_the_page_build_inherits_its_language():
    """``asyncio.ensure_future(_deferred_load())`` from a page: no slot of its
    own, but it copies the page build's context."""
    seen = []

    async def main():
        page = _Page()

        async def deferred():
            await asyncio.sleep(0)
            seen.append(get_language())

        async def build():
            _enter_page(page)
            try:
                set_language('en')
                return asyncio.ensure_future(deferred())
            finally:
                _leave_page()

        task = await asyncio.create_task(build())

        async def other_visitor():
            _enter_page(_Page())
            try:
                set_language('he')
            finally:
                _leave_page()

        await asyncio.create_task(other_visitor())
        await task

    asyncio.run(main())
    assert seen == ['en']


@pytest.mark.parametrize('saved, header, expected', [
    ('en', 'he-IL', 'en'),  # a saved choice wins
    ('he', 'en-US', 'he'),
    (None, 'he-IL,he;q=0.9', 'he'),  # no choice: the browser's language
    (None, 'en-US', 'en'),
    (None, None, 'en'),
])
def test_an_api_request_uses_the_visitors_language(monkeypatch, saved, header, expected):
    """Downloads are built in FastAPI routes: no page, but a request."""
    import web.safe_storage as safe_storage

    storage = {} if saved is None else {'ui_language': saved}
    monkeypatch.setattr(safe_storage, 'safe_user_get',
                        lambda key, default=None: storage.get(key, default))
    headers = {} if header is None else {'accept-language': header}
    request_contextvar.set(SimpleNamespace(headers=headers, state=SimpleNamespace()))
    translations._current_lang = 'he' if expected == 'en' else 'en'  # the process default disagrees
    assert get_language() == expected


def test_an_api_request_resolves_the_language_once(monkeypatch):
    import web.safe_storage as safe_storage

    calls = []
    monkeypatch.setattr(safe_storage, 'safe_user_get',
                        lambda key, default=None: calls.append(key) or 'en')
    request_contextvar.set(SimpleNamespace(headers={}, state=SimpleNamespace()))
    assert [get_language() for _ in range(5)] == ['en'] * 5
    assert calls == ['ui_language']


def test_a_page_render_does_not_change_the_process_default():
    async def main():
        _enter_page(_Page())
        try:
            set_language('en')
        finally:
            _leave_page()

    translations._current_lang = 'he'
    asyncio.run(main())
    assert get_language() == 'he'


def test_a_worker_thread_gets_the_process_default_not_a_page():
    """``run.io_bound`` threads carry no context; they must be handed the
    language. They never pick up a page from another task."""
    seen = []

    async def main():
        _enter_page(_Page())
        try:
            set_language('en')
            worker = threading.Thread(target=lambda: seen.append(get_language()))
            worker.start()
            worker.join()
        finally:
            _leave_page()

    translations._current_lang = 'he'
    asyncio.run(main())
    assert seen == ['he']


def test_using_language_overrides_the_page_and_may_await():
    seen = []

    async def main():
        page = _Page()
        _enter_page(page)
        try:
            set_language('he')
            with using_language('en'):
                await asyncio.sleep(0)
                seen.append(get_language())
            seen.append(get_language())
        finally:
            _leave_page()

    asyncio.run(main())
    assert seen == ['en', 'he']


def test_using_language_in_one_task_is_invisible_to_another():
    seen = []

    async def main():
        inside, release = asyncio.Event(), asyncio.Event()

        async def holder():
            with using_language('en'):
                inside.set()
                await release.wait()

        async def other():
            await inside.wait()
            seen.append(get_language())
            release.set()

        await asyncio.gather(holder(), other())

    translations._current_lang = 'he'
    asyncio.run(main())
    assert seen == ['he']


def test_set_language_in_a_task_outside_a_page_holds_for_that_task(monkeypatch):
    """A request seen earlier in the task (the User simulation's own HTTP calls
    run in the caller's task) must not outrank a language the task set."""
    import web.safe_storage as safe_storage

    monkeypatch.setattr(safe_storage, 'safe_user_get', lambda key, default=None: default)
    seen = []

    async def main():
        request_contextvar.set(SimpleNamespace(headers={}, state=SimpleNamespace()))  # English browser
        set_language('he')
        seen.append(tr(_KEY))

    translations._current_lang = 'en'
    asyncio.run(main())
    assert seen == [_HE]
    assert translations._current_lang == 'en'  # the process default is not touched


@pytest.mark.parametrize('visitor', ['en', 'he'])
def test_the_word_export_names_libraries_in_the_visitors_language(monkeypatch, visitor):
    """The tracker's download case: `/api/export/word` runs in a worker thread
    (a plain ``def`` route) and resolves library names with get_language().
    It used whatever language the last page render left behind."""
    import io
    from unittest.mock import MagicMock

    import docx
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from genizah_core import get_library_display
    from web.api import init_api_routes, state

    bare = FastAPI()

    @bare.middleware('http')
    async def _track_request(request, call_next):  # what NiceGUI's middleware does
        request_contextvar.set(request)
        return await call_next(request)

    init_api_routes(app_override=bare)
    storage = {'ui_language': visitor, 'export_search_payload': {
        'results': [{
            'uid': 'u0',
            'display': {'shelfmark': 'T-S 1.1', 'title': 't', 'id': '9912345678901234',
                        'library_code': 'CUL'},
            'raw_header': 'header_9912345678901201_IE99_P1',
            'snippet': 'a *match*', 'full_text': 'x', 'sort_score': 0.5,
        }],
        'query': 'q', 'mode': 'text', 'gap': None, 'filters': None,
        'warnings': [], 'selected_uids': None,
    }}
    monkeypatch.setattr('web.safe_storage.app',
                        SimpleNamespace(storage=SimpleNamespace(user=storage)))
    mgr = MagicMock()
    mgr.get_meta_for_id.return_value = ('T-S 1.1', 't')
    mgr.get_library_for_id.return_value = 'CUL'
    mgr.parse_full_id_components.return_value = {
        'sys_id': '9912345678901234', 'ie_id': 'IE99', 'p_num': '1', 'fl_id': None}
    other = 'he' if visitor == 'en' else 'en'
    translations._current_lang = other  # what another visitor's render used to leave
    saved_mgr, state.meta_mgr = state.meta_mgr, mgr
    try:
        r = TestClient(bare).get('/api/export/word')
    finally:
        state.meta_mgr = saved_mgr
    assert r.status_code == 200, r.text[:300]
    text = '\n'.join(p.text for p in docx.Document(io.BytesIO(r.content)).paragraphs)
    for table in docx.Document(io.BytesIO(r.content)).tables:
        text += '\n' + '\n'.join(cell.text for row in table.rows for cell in row.cells)
    mine = get_library_display('CUL', short=False, lang=visitor)
    theirs = get_library_display('CUL', short=False, lang=other)
    assert mine != theirs
    assert mine in text and theirs not in text, text[:500]


def test_an_unreadable_request_gives_the_default_not_an_error(monkeypatch):
    """tr() runs everywhere; a request it cannot read must not break the page."""
    import web.safe_storage as safe_storage

    monkeypatch.setattr(safe_storage, 'safe_user_get', lambda key, default=None: default)
    request_contextvar.set(SimpleNamespace())  # no headers, no state
    translations._current_lang = 'he'
    assert tr(_KEY) == _HE
