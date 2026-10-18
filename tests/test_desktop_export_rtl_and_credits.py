# -*- coding: utf-8 -*-
"""Desktop exports: the full MiDRASH citation, and Hebrew Word files that
read right to left from the right margin (owner, 2026-10-17).

* `GenizahGUI._get_credit_header` -- the credit block on the desktop's Word,
  text and CSV exports and the Composition Search workbook -- was a short
  hand-typed copy ("Dataset available at:", no author list). It now carries
  the same lines as the search workbook's Credits and Info sheet
  (`shared.export_dossier.credits_lines`).
* The Composition Search Word export, in Hebrew, came out with right-to-left
  letters flush LEFT: Word reads `jc="right"` on a right-to-left paragraph as
  the end of the line. Driven here through the real `export_comp_report`.
"""
from __future__ import annotations

from types import MethodType, SimpleNamespace

import pytest

import genizah_app
import genizah_core
from genizah_app import GenizahGUI
from shared.export_utils import (DESKTOP_APP_NAME, MIDRASH_CITATION_LINE,
                                 MIDRASH_DATA_SOURCE_LINE, MIDRASH_DATASET_LINE)
from tests.test_comp_flat_markers_and_xlsx_columns import _ExportHost, _patch_dialogs

pytestmark = pytest.mark.gui  # imports genizah_app, which imports PyQt6

docx = pytest.importorskip("docx")


def _lang(monkeypatch, lang):
    monkeypatch.setattr(genizah_core, "CURRENT_LANG", lang)
    monkeypatch.setattr(genizah_app, "CURRENT_LANG", lang, raising=False)


def _credit(local_only=False):
    return GenizahGUI._get_credit_header(SimpleNamespace(), local_only=local_only)


# --- the credit block ---------------------------------------------------------

@pytest.mark.parametrize("lang", ["en", "he"])
def test_the_credit_header_carries_the_full_citation(monkeypatch, lang):
    _lang(monkeypatch, lang)
    lines = _credit().splitlines()
    for line in (MIDRASH_DATA_SOURCE_LINE, MIDRASH_DATASET_LINE, MIDRASH_CITATION_LINE):
        assert line in lines, (lang, lines)
    assert any(DESKTOP_APP_NAME in ln for ln in lines)
    assert any("gershuni@gmail.com" in ln for ln in lines)
    # Still closed by the separator the text and CSV exports print.
    assert "=" * 80 in lines


@pytest.mark.parametrize("lang", ["en", "he"])
def test_a_my_library_only_export_still_omits_the_midrash_lines(monkeypatch, lang):
    _lang(monkeypatch, lang)
    text = _credit(local_only=True)
    assert "MiDRASH" not in text and "zenodo" not in text.lower()
    assert DESKTOP_APP_NAME in text


# --- the Composition Search export, in Hebrew -----------------------------------

class _HebrewHost(_ExportHost):
    """The test window from the xlsx tests, with the REAL credit header and
    the real Word helpers."""

    def __init__(self, main, filtered, ext):
        super().__init__(main, filtered)
        self._ext = ext
        for name in ("_get_credit_header", "_set_table_width_pct",
                     "_add_docx_highlighted_runs", "_set_paragraph_rtl",
                     "_set_table_rtl"):
            setattr(self, name, MethodType(getattr(GenizahGUI, name), self))

    def _default_report_path(self, title, fallback):
        return title + self._ext


def _host(ext):
    page = {"raw_header": "990001_IE1_P000001_FL1", "score": 9,
            "source_ctx": "ברוך *אתה* יי", "text": "ברוך אתה"}
    return _HebrewHost(
        main=[{"type": "manuscript", "sys_id": "990001", "score": 9, "pages": [page]}],
        filtered=[], ext=ext)


def test_the_hebrew_word_export_reads_right_to_left_from_the_right(tmp_path, monkeypatch):
    from docx.oxml.ns import qn
    _lang(monkeypatch, "he")
    out = tmp_path / "report.docx"
    _patch_dialogs(monkeypatch, out)
    _host(".docx").export_comp_report("docx")
    doc = docx.Document(str(out))

    normal_jc = doc.styles["Normal"].element.find(qn("w:pPr"))
    assert normal_jc is None or normal_jc.find(qn("w:jc")) is None, (
        'a "right" Normal style puts every right-to-left paragraph flush left')

    hebrew = [p for p in doc.paragraphs if any("א" <= ch <= "ת" for ch in p.text)]
    assert hebrew
    for p in hebrew:
        ppr = p._p.pPr
        assert ppr is not None and ppr.find(qn("w:bidi")) is not None, p.text
        assert ppr.find(qn("w:jc")) is None, (
            f'right-to-left and aligned "right" (flush left in Word): {p.text!r}')

    texts = [p.text for p in doc.paragraphs]
    assert MIDRASH_CITATION_LINE in texts, texts[:8]
    citation = doc.paragraphs[texts.index(MIDRASH_CITATION_LINE)]
    assert citation._p.pPr.find(qn("w:bidi")) is None
    assert citation.alignment == docx.enum.text.WD_PARAGRAPH_ALIGNMENT.RIGHT

    table = doc.tables[0]
    assert table._tbl.tblPr.find(qn("w:bidiVisual")) is not None
    for cell in table.rows[0].cells:
        if any("א" <= ch <= "ת" for ch in cell.text):
            assert cell.paragraphs[0]._p.pPr.find(qn("w:jc")) is None, cell.text


def test_the_composition_workbook_carries_the_full_citation(tmp_path, monkeypatch):
    openpyxl = pytest.importorskip("openpyxl")
    _lang(monkeypatch, "he")
    out = tmp_path / "report.xlsx"
    _patch_dialogs(monkeypatch, out)
    _host(".xlsx").export_comp_report("xlsx")
    wb = openpyxl.load_workbook(out)
    for title in ("Report View", "Raw Data"):
        ws = next((w for w in wb.worksheets
                   if w.title == genizah_core.tr(title)), None)
        assert ws is not None, [w.title for w in wb.worksheets]
        cells = {str(c.value) for row in ws.iter_rows(max_row=12) for c in row if c.value}
        assert MIDRASH_CITATION_LINE in cells, (title, sorted(cells))
