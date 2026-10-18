# -*- coding: utf-8 -*-
"""Hebrew Word exports: a right-to-left paragraph must not be aligned "right".

Word reads `jc="right"` on a right-to-left (`w:bidi`) paragraph as the END of
the line -- the left margin -- so every Hebrew desktop export came out with
right-to-left letters, flush left (owner, 2026-10-17; rendered through Word
the same day). `shared.docx_export._set_paragraph_rtl` now leaves a Hebrew
paragraph's alignment unset (it then starts at the right margin), keeps a
Latin-only line left-to-right and aligned right, and inserts `w:bidi` /
`w:bidiVisual` where the schema puts them.
"""
from __future__ import annotations

import pytest

docx = pytest.importorskip("docx")

from docx.enum.text import WD_PARAGRAPH_ALIGNMENT  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from shared.docx_export import _set_paragraph_rtl, _set_table_rtl, set_document_rtl  # noqa: E402
from shared.export_utils import MIDRASH_CITATION_LINE  # noqa: E402


def _tags(el):
    return [c.tag for c in el]


def _ppr(p):
    return p._p.pPr


def test_a_hebrew_paragraph_is_right_to_left_with_no_alignment_of_its_own():
    doc = docx.Document()
    p = doc.add_paragraph('סה"כ כתבי יד שנמצאו: 850')
    p.alignment = WD_PARAGRAPH_ALIGNMENT.RIGHT     # what the old helper left
    _set_paragraph_rtl(p)
    assert _ppr(p).find(qn("w:bidi")).get(qn("w:val")) == "1"
    assert _ppr(p).find(qn("w:jc")) is None, (
        'a right-to-left paragraph aligned "right" shows flush LEFT in Word')


def test_a_latin_only_line_stays_left_to_right_and_is_aligned_right():
    doc = docx.Document()
    p = doc.add_paragraph(MIDRASH_CITATION_LINE)
    _set_paragraph_rtl(p)
    assert _ppr(p).find(qn("w:bidi")) is None
    assert p.alignment == WD_PARAGRAPH_ALIGNMENT.RIGHT


def test_a_line_with_no_letters_counts_as_hebrew_text():
    """A separator or a bare count lines up with the Hebrew around it."""
    doc = docx.Document()
    p = doc.add_paragraph("_" * 40)
    _set_paragraph_rtl(p)
    assert _ppr(p).find(qn("w:bidi")) is not None
    assert _ppr(p).find(qn("w:jc")) is None


def test_bidi_sits_before_spacing_where_the_schema_puts_it():
    doc = docx.Document()
    p = doc.add_paragraph("שורה")
    p.paragraph_format.space_after = 0              # creates w:spacing
    _set_paragraph_rtl(p)
    tags = _tags(_ppr(p))
    assert tags.index(qn("w:bidi")) < tags.index(qn("w:spacing")), tags


def test_a_paragraph_already_right_to_left_is_left_as_one():
    doc = docx.Document()
    p = doc.add_paragraph("שורה")
    _set_paragraph_rtl(p)
    _set_paragraph_rtl(p)
    assert len(_ppr(p).findall(qn("w:bidi"))) == 1


def test_a_table_runs_right_to_left_with_bidi_visual_before_its_width():
    doc = docx.Document()
    table = doc.add_table(rows=1, cols=2)
    _set_table_rtl(table)
    tags = _tags(table._tbl.tblPr)
    assert qn("w:bidiVisual") in tags
    assert tags.index(qn("w:bidiVisual")) < tags.index(qn("w:tblW")), tags


def test_the_whole_document_pass_reaches_table_cells():
    doc = docx.Document()
    doc.add_paragraph("כותרת")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "ספרייה"
    table.rows[0].cells[1].text = "T-S 12.34"
    set_document_rtl(doc, tables=(table, None))
    heb, latin = (c.paragraphs[0] for c in table.rows[0].cells)
    assert _ppr(heb).find(qn("w:bidi")) is not None
    assert _ppr(heb).find(qn("w:jc")) is None
    assert _ppr(latin).find(qn("w:bidi")) is None
    assert latin.alignment == WD_PARAGRAPH_ALIGNMENT.RIGHT
    assert _ppr(doc.paragraphs[0]).find(qn("w:bidi")) is not None
